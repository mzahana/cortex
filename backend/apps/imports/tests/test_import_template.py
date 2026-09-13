"""`GET /api/v1/exports/asset-import-template.xlsx`: the blank workbook an
Admin downloads, fills in and uploads back.

Covers the three things that actually matter about it: the header row is
the schema the importer expects (proved by round-tripping a filled copy
through `POST /imports` + commit, not by asserting on strings), uploading
it does NOT touch assets that already exist, and `import.run` gates it.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl.utils import column_index_from_string

from apps.assets.models import Asset
from apps.catalog.models import Category, CustomFieldDef, Location
from apps.common.tests.factories import (
    DEFAULT_TEST_PASSWORD,
    CategoryFactory,
    ProjectFactory,
    TenantFactory,
    UserFactory,
    upgrade_tenant_wide_role,
)
from apps.imports.exports import CORE_EXPORT_COLUMNS
from apps.imports.models import ImportJob
from apps.imports.services import DATA_SHEET_TITLE
from apps.rbac.permission_keys import ROLE_ADMIN, ROLE_MEMBER, ROLE_VIEWER
from apps.tenancy.context import tenant_context

pytestmark = pytest.mark.django_db

TEMPLATE_URL = "/api/v1/exports/asset-import-template.xlsx"


def _login(client, tenant, user):
    response = client.post(
        "/api/v1/auth/login",
        {"tenant": tenant.slug, "email": user.email, "password": DEFAULT_TEST_PASSWORD},
        content_type="application/json",
    )
    assert response.status_code == 200, response.content


def _workbook(response) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(response.content))


def _sheet_text(workbook, title: str) -> str:
    return "\n".join(
        str(cell.value)
        for row in workbook[title].iter_rows()
        for cell in row
        if cell.value is not None
    )


def _lists_columns(workbook) -> dict[str, list[str]]:
    """The `Lists` sheet as `{column heading: [values]}`."""
    rows = list(workbook["Lists"].iter_rows(values_only=True))
    headings = rows[0]
    return {
        heading: [row[index] for row in rows[1:] if row[index] is not None]
        for index, heading in enumerate(headings)
    }


def _header(workbook) -> list[str]:
    sheet = workbook[DATA_SHEET_TITLE]
    return [cell.value for cell in sheet[1]]


def _admin_tenant(with_custom_field: bool = False):
    tenant = TenantFactory()
    admin = UserFactory(tenant=tenant)
    upgrade_tenant_wide_role(admin, ROLE_ADMIN)
    category = CategoryFactory(tenant=tenant, name="Compute")
    if with_custom_field:
        with tenant_context(tenant.id):
            CustomFieldDef.objects.create(
                tenant=tenant,
                category=category,
                key="vram_gb",
                label="VRAM (GB)",
                data_type=CustomFieldDef.DataType.INT,
            )
    return tenant, admin, category


class TestTemplateShape:
    def test_header_row_is_the_import_schema_plus_tenant_custom_fields(self, client):
        tenant, admin, _ = _admin_tenant(with_custom_field=True)
        _login(client, tenant, admin)

        response = client.get(TEMPLATE_URL)
        assert response.status_code == 200, response.content
        assert response["Content-Disposition"].startswith("attachment;")

        assert _header(_workbook(response)) == CORE_EXPORT_COLUMNS + ["vram_gb"]

    def test_data_sheet_is_blank_and_first_so_a_stray_active_tab_cannot_win(self, client):
        tenant, admin, _ = _admin_tenant()
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        assert workbook.sheetnames[0] == DATA_SHEET_TITLE
        # Header row only — nothing that would import as a real asset.
        assert workbook[DATA_SHEET_TITLE].max_row == 1

    def test_custom_field_columns_come_from_config_not_from_existing_assets(self, client):
        """The reason this isn't just "export the CSV and delete the rows":
        a tenant with zero assets still gets every custom-field column."""
        tenant, admin, _ = _admin_tenant(with_custom_field=True)
        _login(client, tenant, admin)
        with tenant_context(tenant.id):
            assert not Asset.objects.exists()

        assert "vram_gb" in _header(_workbook(client.get(TEMPLATE_URL)))

    def test_lists_sheet_carries_the_tenants_own_categories_locations_projects(self, client):
        tenant, admin, _ = _admin_tenant()
        with tenant_context(tenant.id):
            Location.objects.create(tenant=tenant, name="Shelf B2")
        ProjectFactory(tenant=tenant, name="Drone Swarm")
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        columns = _lists_columns(workbook)
        assert columns["Categories"] == ["Compute"]
        assert columns["Locations"] == ["Shelf B2"]
        assert columns["Projects"] == ["Drone Swarm"]
        assert columns["Statuses"] == list(Asset.Status.values)

    def test_instructions_sheet_states_the_add_only_behaviour(self, client):
        tenant, admin, _ = _admin_tenant()
        _login(client, tenant, admin)

        text = _sheet_text(_workbook(client.get(TEMPLATE_URL)), "Instructions")
        assert "never edits, overwrites or deletes" in text

    def test_dropdowns_point_at_the_lists_sheet_for_the_four_closed_columns(self, client):
        tenant, admin, _ = _admin_tenant()
        with tenant_context(tenant.id):
            Location.objects.create(tenant=tenant, name="Shelf B2")
        ProjectFactory(tenant=tenant, name="Drone Swarm")
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        sheet = workbook[DATA_SHEET_TITLE]
        header = [cell.value for cell in sheet[1]]
        validated: dict[str, str] = {}
        for validation in sheet.data_validations.dataValidation:
            letter = str(validation.sqref).split(":")[0].rstrip("0123456789")
            column = header[column_index_from_string(letter) - 1]
            validated[column] = validation.formula1

        assert set(validated) == {"category", "location", "project", "status"}
        assert validated["category"] == "=CortexCategories"
        # The drop-downs must NOT be `errorStyle="stop"` (openpyxl's default),
        # which makes Excel refuse a typed-in value outright — that would
        # make naming a brand-new category from this template impossible,
        # and the create-missing flow unreachable.
        assert {v.errorStyle for v in sheet.data_validations.dataValidation} == {"warning"}
        # Every drop-down must resolve to a real, exactly-sized range on the
        # Lists sheet — an off-by-one here shows up as blank drop-down rows.
        names = {name: dn.attr_text for name, dn in workbook.defined_names.items()}
        assert names["CortexCategories"] == "'Lists'!$A$2:$A$2"
        assert names["CortexStatuses"] == f"'Lists'!$D$2:$D${len(Asset.Status.values) + 1}"

    def test_dropdown_is_omitted_for_a_list_the_tenant_has_not_populated(self, client):
        """A validation over an empty range shows a drop-down of blanks."""
        tenant, admin, _ = _admin_tenant()  # no locations, no projects
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        formulas = {v.formula1 for v in workbook[DATA_SHEET_TITLE].data_validations.dataValidation}
        assert "=CortexCategories" in formulas
        assert "=CortexLocations" not in formulas
        assert "=CortexProjects" not in formulas

    def test_a_duplicated_category_name_is_listed_once(self, client):
        """Two `Battery` categories under different parents are legal in the
        tree; the drop-down must not offer the same string twice."""
        tenant, admin, parent = _admin_tenant()
        with tenant_context(tenant.id):
            Category.objects.create(tenant=tenant, name="Battery")
            Category.objects.create(tenant=tenant, name="Battery", parent=parent)
        _login(client, tenant, admin)

        columns = _lists_columns(_workbook(client.get(TEMPLATE_URL)))
        assert columns["Categories"] == ["Battery", "Compute"]


class TestRoundTrip:
    def test_filled_template_imports_and_leaves_existing_assets_untouched(
        self, client, django_capture_on_commit_callbacks
    ):
        tenant, admin, category = _admin_tenant(with_custom_field=True)
        with tenant_context(tenant.id):
            existing = Asset.objects.create(
                tenant=tenant, category=category, name="Already here", condition="Mint"
            )
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        sheet = workbook[DATA_SHEET_TITLE]
        header = [cell.value for cell in sheet[1]]
        row = {"name": "Jetson Orin #3", "category": "Compute", "vram_gb": 64}
        sheet.append([row.get(column) for column in header])
        buffer = io.BytesIO()
        workbook.save(buffer)

        with django_capture_on_commit_callbacks(execute=True):
            upload = client.post(
                "/api/v1/imports",
                {
                    "file": SimpleUploadedFile(
                        "filled.xlsx",
                        buffer.getvalue(),
                        content_type=(
                            "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
                        ),
                    )
                },
            )
        assert upload.status_code == 202, upload.content
        import_id = upload.json()["id"]

        detail = client.get(f"/api/v1/imports/{import_id}").json()
        assert detail["status"] == "dry_run_succeeded", detail
        # Every template header auto-maps: nothing fell through to `ignore`.
        assert detail["report"]["invalid_count"] == 0
        assert detail["report"]["total_rows"] == 1

        with django_capture_on_commit_callbacks(execute=True):
            commit = client.post(
                f"/api/v1/imports/{import_id}/commit", {}, content_type="application/json"
            )
        assert commit.status_code == 202, commit.content
        assert (
            client.get(f"/api/v1/imports/{import_id}").json()["status"]
            == ImportJob.Status.COMMITTED
        )

        with tenant_context(tenant.id):
            names = sorted(Asset.objects.values_list("name", flat=True))
            assert names == ["Already here", "Jetson Orin #3"]
            existing.refresh_from_db()
            assert existing.condition == "Mint"
            imported = Asset.objects.get(name="Jetson Orin #3")
            assert [(v.field_def.key, v.value) for v in imported.field_values.all()] == [
                ("vram_gb", 64)
            ]

    def test_instructions_sheet_is_never_imported_even_when_left_active(
        self, client, django_capture_on_commit_callbacks
    ):
        """Excel/Sheets save whichever tab was selected as the active one —
        the parser must still read the `Assets` sheet."""
        tenant, admin, _ = _admin_tenant()
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        sheet = workbook[DATA_SHEET_TITLE]
        header = [cell.value for cell in sheet[1]]
        sheet.append(
            [{"name": "Bench PSU", "category": "Compute"}.get(column) for column in header]
        )
        workbook.active = workbook.sheetnames.index("Instructions")
        buffer = io.BytesIO()
        workbook.save(buffer)

        with django_capture_on_commit_callbacks(execute=True):
            upload = client.post(
                "/api/v1/imports",
                {"file": SimpleUploadedFile("filled.xlsx", buffer.getvalue())},
            )
        assert upload.status_code == 202, upload.content
        report = client.get(f"/api/v1/imports/{upload.json()['id']}").json()["report"]
        assert report["total_rows"] == 1
        assert report["rows"][0]["values"]["name"] == "Bench PSU"


class TestRbacAndTenancy:
    @pytest.mark.parametrize("role", [ROLE_MEMBER, ROLE_VIEWER])
    def test_non_admin_cannot_download_the_template(self, client, role):
        tenant = TenantFactory()
        user = UserFactory(tenant=tenant)
        upgrade_tenant_wide_role(user, role)
        _login(client, tenant, user)

        assert client.get(TEMPLATE_URL).status_code == 403

    def test_anonymous_is_rejected(self, client):
        assert client.get(TEMPLATE_URL).status_code in (401, 403)

    def test_template_never_lists_another_tenants_categories(self, client):
        tenant, admin, _ = _admin_tenant()
        other = TenantFactory()
        CategoryFactory(tenant=other, name="SecretOtherTenantCategory")
        _login(client, tenant, admin)

        workbook = _workbook(client.get(TEMPLATE_URL))
        assert "SecretOtherTenantCategory" not in _sheet_text(workbook, "Lists")
        assert "SecretOtherTenantCategory" not in _sheet_text(workbook, "Instructions")
